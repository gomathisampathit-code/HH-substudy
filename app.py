import streamlit as st
import pandas as pd
import numpy as np
from datetime import datetime
from openpyxl.styles import Alignment, Border, Side, Font
from io import BytesIO
from openpyxl import Workbook


# =========================================================
# PASSWORD PROTECTION
# =========================================================

st.title("🔒 RespIndNet HH Substudy Specimen Transfer Form (Virology)")

password = st.text_input("Enter Password:", type="password")

if password != "HH123":
    st.warning("Please enter the correct password.")
    st.stop()


# =========================================================
# LOAD GOOGLE SHEET
# =========================================================

# HH Sample Collection Form
sheet_id = "1wZNK_uRuTlFWtS4HAS-5dvxfMZzhe0u4vwgBWkvKD-4"

csv_url = (
    f"https://docs.google.com/spreadsheets/d/"
    f"{sheet_id}/gviz/tq?tqx=out:csv"
)

df = pd.read_csv(csv_url, on_bad_lines="skip")


# =========================================================
# NORMALIZE COLUMNS
# =========================================================

df.columns = df.columns.str.strip().str.lower()

df["submissiondate"] = pd.to_datetime(
    df["submissiondate"],
    errors="coerce"
)

df["submissiondate"] = df["submissiondate"].dt.tz_localize(None)

df["dt_sample"] = pd.to_datetime(
    df["dt_sample"],
    errors="coerce"
)

df["dt_sample"] = df["dt_sample"].dt.tz_localize(None)


# =========================================================
# FILTER TODAY'S COLLECTED SAMPLES
# =========================================================

df["sample_collected"] = (
    df["sample_collected"]
    .astype(str)
    .str.strip()
    .str.lower()
)

today_str = pd.Timestamp.today().strftime("%Y-%m-%d")

df_today = df[
    (df["submissiondate"].dt.strftime("%Y-%m-%d") == today_str)
    & (df["sample_collected"] == "yes")
].copy()

df_today = df_today.reset_index(drop=True)


# =========================================================
# SAMPLE TYPE
# =========================================================

# HH substudy samples are Respiratory swabs
df_today["type_of_sample"] = "Respiratory swab"


# =========================================================
# BARCODE ID
# =========================================================

df_today["barcode_id"] = df_today["sample_id"]


# =========================================================
# SAMPLE SEQUENCE PER MEMBER
# =========================================================

df_today["sample_sequence"] = (
    df_today.groupby("member_id").cumcount() + 1
)


# =========================================================
# SPLIT MEMBER ID AND NAME
# =========================================================

def split_member(val):

    if pd.isna(val) or str(val).strip() == "":
        return "", ""

    parts = str(val).split(",", 1)

    if len(parts) == 2:
        return parts[0].strip(), parts[1].strip()

    return parts[0].strip(), ""


df_today["member_id_clean"], df_today["member_name"] = zip(
    *df_today["member_id"].map(split_member)
)


# =========================================================
# INDEX CASE
# =========================================================

df_today["index_case_label"] = (
    pd.to_numeric(
        df_today["index_case"],
        errors="coerce"
    )
    .map({
        1: "Index",
        2: "Contact"
    })
    .fillna("")
)

name_lookup = {}

for _, row in df_today.iterrows():

    member_id = str(row["member_id_clean"]).strip()
    member_name = str(row["member_name"]).strip()

    # Ignore empty / NaN names
    if (
        member_name
        and member_name.lower() != "nan"
        and member_name.lower() != "none"
    ):

        # IDs ending with -2 are used to identify the mother
        if member_id.endswith("-2"):

            # Remove the final "-2"
            base_id = member_id[:-2]

            # Store mother name against base ID
            name_lookup[base_id] = member_name


# =========================================================
# CREATE FINAL DISPLAY NAME
# =========================================================

def get_display_name(row):

    member_id = str(row["member_id_clean"]).strip()
    member_name = str(row["member_name"]).strip()

    # -----------------------------------------------------
    # If name is already available, use it
    # -----------------------------------------------------

    if (
        member_name
        and member_name.lower() != "nan"
        and member_name.lower() != "none"
    ):
        return member_name

    # -----------------------------------------------------
    # If name is blank, find corresponding -2 member
    # -----------------------------------------------------

    if member_id in name_lookup:

        mother_name = name_lookup[member_id]

        return f"{mother_name}'s baby"

    # -----------------------------------------------------
    # If no matching name is found, leave blank
    # -----------------------------------------------------

    return ""


df_today["display_name"] = df_today.apply(
    get_display_name,
    axis=1
)


# =========================================================
# FINAL TABLE
# =========================================================

table = pd.DataFrame({

    "S.NO": range(
        1,
        len(df_today) + 1
    ),

    "BARCODE ID":
        df_today["barcode_id"],

    "Episode":
        df_today["episode1"],

    "Index case":
        df_today["index_case_label"],

    "S.TYPE":
        df_today["type_of_sample"],

    "S.PER IND":
        df_today["sample_sequence"],

    "HH substudy member ID":
        df_today["member_id_clean"],

    "NAME":
        df_today["display_name"],

    "Day":
        df_today.get(
            "sample_timepoint",
            ""
        ),

    "S.C DATE/TIME":
        df_today["dt_sample"],

    "STUDY":
        "",

    "RECEIVED BY":
        "",

    "VOL (VIRO)":
        "",

    "REMARKS(LYSED/LIPEMIC/ICTERIC/SAMPLE SPILLAGE)":
        ""
})


# =========================================================
# DISPLAY TABLE
# =========================================================

st.subheader("📋 Generated Table")

st.dataframe(
    table,
    use_container_width=True
)


# =========================================================
# DOWNLOAD EXCEL
# =========================================================

if len(table) > 0:

    today_str_display = datetime.today().strftime(
        "%d-%m-%Y"
    )

    excel_filename = (
        f"{today_str_display}_"
        f"RespIndNet_HH_STF(Field_to_Virology).xlsx"
    )


    # -----------------------------------------------------
    # CREATE WORKBOOK
    # -----------------------------------------------------

    wb = Workbook()

    ws = wb.active

    ws.title = "Report"


    # -----------------------------------------------------
    # BORDER
    # -----------------------------------------------------

    thin = Side(
        border_style="thin",
        color="000000"
    )

    border = Border(
        top=thin,
        left=thin,
        right=thin,
        bottom=thin
    )


    # =====================================================
    # HEADER ROW 1
    # =====================================================

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=table.shape[1]
    )

    cell = ws.cell(
        row=1,
        column=1,
        value="RespIndNet_HH substudy Specimen Transfer Form (Virology)"
    )

    cell.font = Font(bold=True)

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


    # Apply border to merged header
    for col in range(
        1,
        table.shape[1] + 1
    ):

        ws.cell(
            row=1,
            column=col
        ).border = border


    # =====================================================
    # EXTRA SPAN ROW 2
    # =====================================================

    spans = [

        "Sample Shipment Date and Time:",

        "Field Manager Sign/Initials:",

        "Virology Staff Sign/Initials:",

        "To be filled by Virology"

    ]


    # Divide columns into 4 sections
    base_width = table.shape[1] // 4

    col_split = [

        base_width,

        base_width,

        table.shape[1] - (
            3 * base_width
        ),

        base_width

    ]


    start_col = 1


    for i, val in enumerate(spans):

        end_col = (
            start_col
            + col_split[i]
            - 1
        )

        ws.merge_cells(
            start_row=2,
            start_column=start_col,
            end_row=2,
            end_column=end_col
        )

        cell = ws.cell(
            row=2,
            column=start_col,
            value=val
        )

        cell.font = Font(bold=True)

        cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True
        )


        # Apply borders
        for col in range(
            start_col,
            end_col + 1
        ):

            ws.cell(
                row=2,
                column=col
            ).border = border


        start_col = end_col + 1


    # =====================================================
    # COLUMN HEADERS - ROW 3
    # =====================================================

    for j, col_name in enumerate(
        table.columns,
        1
    ):

        c = ws.cell(
            row=3,
            column=j,
            value=col_name
        )

        c.font = Font(bold=True)

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        c.border = border


    # =====================================================
    # DATA ROWS
    # =====================================================

    for i, row in table.iterrows():

        for j, val in enumerate(
            row,
            1
        ):

            # Convert pandas NaN to blank
            if pd.isna(val):
                val = ""

            c = ws.cell(
                row=i + 4,
                column=j,
                value=val
            )

            c.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True
            )

            c.border = border


    # =====================================================
    # SAVE TO MEMORY
    # =====================================================

    buffer = BytesIO()

    wb.save(buffer)

    buffer.seek(0)


    # =====================================================
    # DOWNLOAD BUTTON
    # =====================================================

    st.download_button(

        label="⬇️ Download Excel",

        data=buffer,

        file_name=excel_filename,

        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )
